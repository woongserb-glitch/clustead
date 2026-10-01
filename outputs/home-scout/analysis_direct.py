"""Read-only MOLIT direct-trade side index; never rebuild the transaction master.

Reuse ``trade_identity(master_row)`` and ``load_direct_keys()`` in cohort analysis.
The key includes exact raw apartment name, normalized road, lot, date, area,
floor and amount, so same-name complexes cannot leak exclusions to each other.
"""
import csv
import json
import sys
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from transaction_layer_utils import clean_text, normalize_address, normalize_dong, normalize_lot_part
from build_molit_transaction_master import normalize_row, is_cancelled_contract

DEFAULT_OUTPUT = Path(__file__).with_name("analysis_direct_keys.json")
KEY_FIELDS = ("gu", "dong", "apartment_name", "normalized_road_address", "bonbun", "bubun",
              "contract_date", "area_m2", "floor", "trade_price_manwon")


def number_key(value):
    text = str(value or "").strip().replace(",", "")
    if not text:
        return ""
    try:
        return format(Decimal(text).normalize(), "f")
    except InvalidOperation:
        return text


def trade_identity(row):
    """Stable key for raw-normalized or CSV-master trade rows (no approximate match)."""
    main = normalize_lot_part(row.get("bonbun"))
    sub = normalize_lot_part(row.get("bubun"))
    if not main:
        parts = clean_text(row.get("jibun")).split("-", 1)
        main = normalize_lot_part(parts[0])
        sub = normalize_lot_part(parts[1] if len(parts) > 1 else "")
    return (
        clean_text(row.get("gu")), normalize_dong(row.get("dong")),
        clean_text(row.get("apartment_name") or row.get("apt_name_raw")),
        normalize_address(row.get("normalized_road_address") or row.get("road_address")),
        main, "" if sub == "0" else sub,
        clean_text(row.get("contract_date") or row.get("deal_date")),
        number_key(row.get("area_m2")), number_key(row.get("floor")),
        number_key(row.get("trade_price_manwon") or row.get("deal_amount")),
    )


def load_direct_keys(path=DEFAULT_OUTPUT):
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    return {tuple(key) for key in payload["direct_keys_in_master"]}


def raw_rows(path):
    """Stream the spreadsheet without saving it or calling pipeline builders."""
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[book.sheetnames[0]]
        if sheet.max_row == 1:
            sheet.reset_dimensions()
        headers = None
        for raw in sheet.iter_rows(values_only=True):
            values = [str(value).strip() if value is not None else "" for value in raw]
            if headers is None:
                compact = ["".join(value.split()) for value in values]
                if all(field in compact for field in ("시군구", "단지명", "계약년월")):
                    headers = compact
                continue
            row = dict(zip(headers, values))
            if row.get("단지명"):
                yield row
    finally:
        book.close()


def main():
    direct, cancelled = set(), set()
    raw_counts = {}
    for filename in ("trade_2025.xlsx", "trade_2026.xlsx"):
        counts = Counter()
        for row in raw_rows(ROOT / "data/transactions/raw/molit" / filename):
            counts["rows"] += 1
            is_direct = clean_text(row.get("거래유형")) == "직거래"
            is_cancelled = is_cancelled_contract(row)
            counts["direct_rows"] += is_direct
            counts["cancelled_rows"] += is_cancelled
            if not is_direct and not is_cancelled:
                continue
            normalized = normalize_row(row, filename, "trade")
            if normalized is None:
                counts["normalization_excluded_rows"] += 1
                continue
            key = trade_identity(normalized)
            if is_direct:
                direct.add(key)
            if is_cancelled:
                cancelled.add(key)
        raw_counts[filename] = dict(counts)
        print(filename, dict(counts), flush=True)

    matched, cancelled_matched = set(), set()
    direct_by_month = Counter()
    master_stats = Counter()
    with (ROOT / "data/transactions/transaction_master.csv").open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if row["transaction_type"] != "trade":
                continue
            master_stats["trade_rows"] += 1
            key = trade_identity(row)
            if key in direct:
                matched.add(key)
                direct_by_month[row["contract_date"][:7]] += 1
                master_stats["direct_trade_rows"] += 1
            if key in cancelled:
                cancelled_matched.add(key)
                master_stats["cancelled_exact_rows"] += 1

    payload = {"key_fields": KEY_FIELDS,
               "definition": "raw 거래유형=직거래, exact identity joined to current master; no fuzzy/name-only fallback",
               "sources": raw_counts, "master_counts": dict(master_stats),
               "raw_direct_unique_keys": len(direct), "raw_cancelled_unique_keys": len(cancelled),
               "direct_unique_keys_in_master": len(matched), "cancelled_unique_keys_in_master": len(cancelled_matched),
               "direct_rows_by_month": dict(sorted(direct_by_month.items())),
               "direct_keys_in_master": sorted(matched), "cancelled_keys_in_master": sorted(cancelled_matched)}
    DEFAULT_OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in payload.items() if not key.endswith("keys_in_master")}, ensure_ascii=False, indent=2))
    print("direct_unique_keys_in_master",len(matched),"cancelled_unique_keys_in_master",len(cancelled_matched))
    print("wrote", DEFAULT_OUTPUT)


if __name__ == "__main__":
    main()
