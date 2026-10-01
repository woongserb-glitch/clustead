"""Read-only transaction profile. Outputs only an analysis JSON beside this script."""
import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from transaction_layer_utils import normalize_name, normalize_address, normalize_dong
from build_transaction_summary import trusted_mapping


def read(path, encoding="utf-8-sig"):
    with path.open(encoding=encoding, newline="") as handle:
        yield from csv.DictReader(handle)


def main():
    path = ROOT / "data/transactions/transaction_master.csv"
    mapping = list(read(ROOT / "data/transactions/apartment_transaction_mapping.csv"))
    monthly = defaultdict(Counter)
    kinds = Counter()
    dates = []
    groups = Counter()
    lots = defaultdict(Counter)
    raw_names = defaultdict(set)
    example_rows = []
    invalid = Counter()
    aliases = Counter()
    master_trade_keys = set()
    for row in read(path):
        kind = row["transaction_type"]
        month = row["contract_date"][:7]
        label = kind
        if kind == "rent":
            value = row["monthly_rent_manwon"]
            label = "jeonse" if value and float(value) == 0 else "monthly" if value and float(value) > 0 else "unknown_rent"
        kinds[label] += 1
        monthly[month][label] += 1
        dates.append(row["contract_date"]) if len(dates) < 1 else None
        key = (row["gu"].strip(), normalize_dong(row["dong"]), normalize_name(row["apartment_name"] or row["apt_name_raw"]), normalize_address(row["normalized_road_address"] or row["road_address"]))
        groups[key] += 1
        lots[key][row["jibun"]] += 1
        raw_names[key].add(row["apartment_name"])
        for field in ("contract_date", "area_m2"):
            if not row[field]:
                invalid[field] += 1
        for first, second in (("contract_date", "deal_date"), ("trade_price_manwon", "deal_amount"), ("deposit_manwon", "deposit_amount"), ("monthly_rent_manwon", "monthly_rent")):
            if row[first] != row[second]:
                aliases[first + "!=" + second] += 1
        if row["dong"] == "오류동" and "금강" in row["apartment_name"]:
            example_rows.append({key: row[key] for key in ("gu", "dong", "apartment_name", "road_address", "jibun", "transaction_type", "contract_date", "area_m2", "trade_price_manwon", "deposit_manwon", "monthly_rent_manwon", "floor")})
        if kind == "trade":
            master_trade_keys.add((row["gu"],row["dong"],row["apartment_name"],row["contract_date"],row["area_m2"],row["floor"],row["trade_price_manwon"]))
    by_triple = defaultdict(list)
    by_gu_name = defaultdict(list)
    for key in groups:
        by_triple[key[:3]].append(key)
        by_gu_name[(key[0],key[2])].append(key)
    ambiguity = []
    for key, groupkeys in by_triple.items():
        if len(groupkeys) > 1:
            ambiguity.append({"key": list(key), "rows": sum(groups[k] for k in groupkeys),
                              "sites": [{"road":k[3],"lots":dict(lots[k]),"names":sorted(raw_names[k]),"rows":groups[k]} for k in groupkeys]})
    mapping_stats = Counter()
    selected_by_group = defaultdict(list)
    mapping_examples = []
    for m in mapping:
        identity=(m["livefit_name"],m["gu"],m["dong"])
        if not trusted_mapping(m) or not normalize_name(m["transaction_apt_name"]):
            mapping_stats["untrusted_or_no_name"] += 1
            continue
        key=(m["gu"].strip(),normalize_dong(m["dong"]),normalize_name(m["transaction_apt_name"]),normalize_address(m["transaction_road_address"]))
        selected=[]
        if key in groups:
            mode="exact_road"
            selected=[key]
        elif key[:3] in by_triple:
            mode="same_dong_name_fallback"
            selected=by_triple[key[:3]]
        else:
            candidates=by_gu_name.get((key[0],key[2]),[])
            if len({k[1] for k in candidates})==1:
                mode="single_other_dong_fallback"
                selected=candidates
            else:
                mode="no_or_ambiguous_match"
        mapping_stats[mode]+=1
        for selected_key in selected:
            selected_by_group[selected_key].append(identity)
        if mode!="exact_road" or (m["dong"]=="오류동" and "금강" in m["transaction_apt_name"]):
            mapping_examples.append({"identity":identity,"mode":mode,"mapping":m,"matched_group_count":len(selected),"rows":sum(groups[k] for k in selected)})
    overlaps=[{"transaction_key":key,"rows":groups[key],"clustead_keys":identities} for key,identities in selected_by_group.items() if len(set(identities))>1]
    raw_profiles=[]
    direct_master=set()
    import openpyxl
    from build_molit_transaction_master import normalize_row, is_cancelled_contract
    for filename in ("trade_2025.xlsx","trade_2026.xlsx","rent_2026.xlsx"):
        book=openpyxl.load_workbook(ROOT / "data/transactions/raw/molit" / filename, read_only=True, data_only=True)
        sheet=book[book.sheetnames[0]]
        if sheet.max_row==1:
            sheet.reset_dimensions()
        headers=None
        counts=Counter()
        for raw in sheet.iter_rows(values_only=True):
            values=[str(x).strip() if x is not None else "" for x in raw]
            if headers is None:
                compact=["".join(x.split()) for x in values]
                if all(x in compact for x in ("시군구","단지명","계약년월")):
                    headers=compact
                    if filename.startswith("rent"):
                        break
                continue
            row=dict(zip(headers,values))
            if not row.get("단지명"):
                continue
            counts["rows"]+=1
            deal_kind=row.get("거래유형", "")
            counts["transaction_mode:"+deal_kind]+=1
            if is_cancelled_contract(row):
                counts["cancelled_rows"]+=1
            if deal_kind=="직거래":
                parsed=normalize_row(row,filename,"trade")
                if parsed:
                    key=tuple(str(parsed[field]) for field in ("gu","dong","apartment_name","contract_date","area_m2","floor","trade_price_manwon"))
                    if key in master_trade_keys:
                        direct_master.add(key)
        book.close()
        raw_profiles.append({"file":filename,"headers":headers,"counts":dict(counts)})
    output={"master_schema":list(next(read(path))),"rows":sum(kinds.values()),"types":dict(kinds),"monthly":dict(sorted(monthly.items())),
            "invalid":dict(invalid),"alias_disagreement":dict(aliases),"mapping_rows":len(mapping),"mapping_modes":dict(mapping_stats),
            "multiple_road_same_name_groups":len(ambiguity),"ambiguous_name_rows":sum(x["rows"] for x in ambiguity),
            "mapping_overlaps":overlaps,"mapping_nonexact_examples":mapping_examples,
            "oryudong_geumgang_rows":example_rows,"oryudong_geumgang_sites":[x for x in ambiguity if x["key"][:2]==["구로구","오류동"] and "금강" in x["key"][2]],
            "raw_profiles":raw_profiles,"direct_trade_keys_in_master":len(direct_master)}
    target=Path(__file__).with_name("analysis_profile.json")
    target.write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({key:output[key] for key in ("rows","types","monthly","invalid","alias_disagreement","mapping_rows","mapping_modes","multiple_road_same_name_groups","ambiguous_name_rows","direct_trade_keys_in_master")},ensure_ascii=False,indent=2))
    print("mapping_overlaps",len(overlaps))
    print("raw_profiles",json.dumps(raw_profiles,ensure_ascii=False))
    print("wrote",target)


if __name__=="__main__":
    main()
